# -*- coding: utf-8 -*-
"""Genera josh_auditoria_manual_2025.xlsx a partir de josh_excel_auditoria."""
from datetime import date

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from josh_excel_auditoria import construye

DIAS = ["lunes", "martes", "miércoles", "jueves", "viernes"]

AR = "Arial"
F_TIT = Font(name=AR, size=13, bold=True)
F_CAB = Font(name=AR, size=10, bold=True, color="FFFFFF")
F_TXT = Font(name=AR, size=10)
F_LEY = Font(name=AR, size=9, italic=True, color="666666")
F_EJ  = Font(name=AR, size=9, italic=True, color="888888")
AMARILLO = PatternFill("solid", fgColor="FFF2A8")
GRIS     = PatternFill("solid", fgColor="EFEFEF")
CAB_LDN  = PatternFill("solid", fgColor="1F4E79")
CAB_NY   = PatternFill("solid", fgColor="7B3F00")
CAB_GEN  = PatternFill("solid", fgColor="404040")
BORDE = Border(*[Side(style="thin", color="C0C0C0")] * 4)
CENTRO = Alignment(horizontal="center")

filas = construye(date(2025, 1, 1), date(2025, 12, 31))

wb = Workbook()
ws = wb.active
ws.title = "Auditoria 2025"

ws["A1"] = "Auditoría manual — Modelos de Josh, EURUSD 2025"
ws["A1"].font = F_TIT
ws["A2"] = ("Modelos generados con la lógica del indicador JoshModels_H1: un modelo máximo por sesión "
            "(Londres: LR2.1 > LR1 > A1 > LR2.2 · NY: NYR2 > NYC1 > NYR1), rango de Asia 17:00–00:00 NY, "
            "todas las señales exigen confirmación (cierre de vuelta dentro), riesgo mínimo 5 pips, sin OSOK, sin NYC2, sin News Protocol. Horas en NY.")
ws["A2"].font = F_LEY
ws["A3"] = ("Rellena SOLO las celdas amarillas: OK? (desplegable OK/NO) y, si es NO, la explicación. "
            "Si el modelo es «—» y ese día sí había un modelo válido, escribe cuál en la explicación.")
ws["A3"].font = Font(name=AR, size=9, italic=True, bold=True, color="9C6500")

# fila de ejemplo
ws["A5"] = "EJEMPLO"
ws["A5"].font = F_EJ
ej = ["", "", "LR1", "corto", "03:00", 1.03500, "NO", "el raid válido era del PDH, no del AH",
      "NYC1", "corto", "08:00", 1.03900, "OK", ""]
for i, v in enumerate(ej[2:], start=3):
    c = ws.cell(row=5, column=i, value=v)
    c.font = F_EJ
    if i in (6, 12):
        c.number_format = "0.00000"

CAB = ["Fecha", "Día", "Modelo", "Dir", "Entrada", "SL", "OK?", "Explicación (si NO)",
       "Modelo", "Dir", "Entrada", "SL", "OK?", "Explicación (si NO)"]
FILA_CAB = 6
for i, t in enumerate(CAB, start=1):
    c = ws.cell(row=FILA_CAB, column=i, value=t)
    c.font = F_CAB
    c.fill = CAB_GEN if i <= 2 else (CAB_LDN if i <= 8 else CAB_NY)
    c.alignment = CENTRO
    c.border = BORDE
ws.cell(row=FILA_CAB - 1, column=3, value="SESIÓN LONDRES").font = Font(name=AR, size=10, bold=True, color="1F4E79")
ws.cell(row=FILA_CAB - 1, column=9, value="SESIÓN NUEVA YORK").font = Font(name=AR, size=10, bold=True, color="7B3F00")

fila = FILA_CAB + 1
primera_dato = fila
for f in filas:
    ws.cell(row=fila, column=1, value=f["fecha"]).number_format = "DD/MM/YYYY"
    ws.cell(row=fila, column=2, value=DIAS[f["fecha"].weekday()])
    if f["nota"] == "sin sesion":
        ws.cell(row=fila, column=3, value="sin sesión")
        ws.cell(row=fila, column=9, value="sin sesión")
        for col in range(1, 15):
            c = ws.cell(row=fila, column=col)
            c.font = F_TXT
            c.fill = GRIS
            c.border = BORDE
        fila += 1
        continue
    for base, clave in ((3, "ldn"), (9, "ny")):
        s = f[clave]
        if s is None:
            ws.cell(row=fila, column=base, value="—")
        else:
            ws.cell(row=fila, column=base, value=s["modelo"])
            ws.cell(row=fila, column=base + 1, value="largo" if s["dir"] == 1 else "corto")
            ws.cell(row=fila, column=base + 2, value=s["entrada"])
            c = ws.cell(row=fila, column=base + 3, value=round(s["sl"], 5))
            c.number_format = "0.00000"
    for col in range(1, 15):
        c = ws.cell(row=fila, column=col)
        c.font = F_TXT
        c.border = BORDE
        if col in (4, 5, 7, 10, 11, 13):
            c.alignment = CENTRO
        if col in (7, 8, 13, 14):
            c.fill = AMARILLO
    fila += 1
ultima_dato = fila - 1

# desplegable OK/NO
dv = DataValidation(type="list", formula1='"OK,NO"', allow_blank=True)
ws.add_data_validation(dv)
dv.add(f"G{primera_dato}:G{ultima_dato}")
dv.add(f"M{primera_dato}:M{ultima_dato}")

# resumen con formulas (se rellena solo segun el usuario marca OK/NO)
rs = ultima_dato + 2
ws.cell(row=rs, column=1, value="RESUMEN (automático)").font = Font(name=AR, size=10, bold=True)
ws.cell(row=rs, column=3, value="(las cifras aparecen al abrir el fichero en Excel o Google Sheets)").font = F_LEY
resumen = [
    ("Señales Londres",  f'=COUNTA(C{primera_dato}:C{ultima_dato})-COUNTIF(C{primera_dato}:C{ultima_dato},"—")-COUNTIF(C{primera_dato}:C{ultima_dato},"sin sesión")'),
    ("  · OK",           f'=COUNTIF(G{primera_dato}:G{ultima_dato},"OK")'),
    ("  · NO",           f'=COUNTIF(G{primera_dato}:G{ultima_dato},"NO")'),
    ("Señales NY",       f'=COUNTA(I{primera_dato}:I{ultima_dato})-COUNTIF(I{primera_dato}:I{ultima_dato},"—")-COUNTIF(I{primera_dato}:I{ultima_dato},"sin sesión")'),
    ("  · OK",           f'=COUNTIF(M{primera_dato}:M{ultima_dato},"OK")'),
    ("  · NO",           f'=COUNTIF(M{primera_dato}:M{ultima_dato},"NO")'),
]
for k, (etq, formula) in enumerate(resumen):
    ws.cell(row=rs + 1 + k, column=1, value=etq).font = F_TXT
    c = ws.cell(row=rs + 1 + k, column=2, value=formula)
    c.font = F_TXT

anchos = {"A": 11, "B": 10, "C": 9, "D": 7, "E": 13, "F": 10, "G": 6, "H": 34,
          "I": 9, "J": 7, "K": 13, "L": 10, "M": 6, "N": 34}
for col, w in anchos.items():
    ws.column_dimensions[col].width = w
ws.freeze_panes = f"A{FILA_CAB + 1}"

ws["A1"].comment = Comment(
    "Generado desde los datos M1 de HistData del repo (hora NY con DST). "
    "Logica identica al indicador JoshModels_H1.pine en su configuracion por defecto. "
    "SL segun la definicion de cada modelo del spec.", "generador")

wb.save("josh_auditoria_manual_2025.xlsx")
print("guardado josh_auditoria_manual_2025.xlsx, filas de datos:", ultima_dato - primera_dato + 1)
